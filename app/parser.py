"""
Módulo de parsing de documentos.

Responsável por extrair texto bruto de arquivos PDF, DOCX e TXT,
delegando para a biblioteca apropriada conforme a extensão.

Estratégia para PDFs:
    1. Tentativa primária com pdfplumber (mais robusto em layouts complexos)
    2. Fallback para PyPDF2 caso pdfplumber falhe
"""

from pathlib import Path
from typing import Union

import pdfplumber
from PyPDF2 import PdfReader
from docx import Document

from app.utils import clean_text, get_file_extension, setup_logger

logger = setup_logger(__name__)


class UnsupportedFileError(Exception):
    """Erro lançado quando o tipo de arquivo não é suportado."""


class DocumentParseError(Exception):
    """Erro genérico durante o parsing de um documento."""


class DocumentParser:
    """
    Classe responsável por extrair texto bruto de documentos.

    Suporta: PDF (.pdf), Word (.docx) e texto plano (.txt).
    """

    SUPPORTED_EXTENSIONS = {"pdf", "docx", "txt"}

    def __init__(self, file_path: Union[str, Path]):
        """
        Inicializa o parser para um arquivo específico.

        Args:
            file_path: Caminho para o arquivo a ser parseado.

        Raises:
            FileNotFoundError: Se o arquivo não existir.
            UnsupportedFileError: Se a extensão não for suportada.
        """
        self.file_path = Path(file_path)

        if not self.file_path.exists():
            raise FileNotFoundError(f"Arquivo não encontrado: {self.file_path}")

        self.extension = get_file_extension(str(self.file_path))

        if self.extension not in self.SUPPORTED_EXTENSIONS:
            raise UnsupportedFileError(
                f"Extensão '.{self.extension}' não suportada. "
                f"Suportadas: {sorted(self.SUPPORTED_EXTENSIONS)}"
            )

        logger.info(f"Parser inicializado para {self.file_path.name} (.{self.extension})")

    def parse(self) -> str:
        """
        Extrai o texto do documento, delegando para o método apropriado.

        Returns:
            Texto extraído, limpo e normalizado.

        Raises:
            DocumentParseError: Se a extração falhar.
        """
        try:
            if self.extension == "pdf":
                raw_text = self._parse_pdf()
            elif self.extension == "docx":
                raw_text = self._parse_docx()
            elif self.extension == "txt":
                raw_text = self._parse_txt()
            else:
                # Já validado no __init__, mas mantém safety net
                raise UnsupportedFileError(f"Extensão '{self.extension}' não suportada")

            cleaned = clean_text(raw_text)
            logger.info(
                f"Documento parseado com sucesso: {len(cleaned)} caracteres extraídos"
            )
            return cleaned

        except (UnsupportedFileError, FileNotFoundError):
            raise
        except Exception as e:
            logger.exception(f"Erro ao parsear documento: {e}")
            raise DocumentParseError(f"Falha ao parsear {self.file_path.name}: {e}") from e

    def _parse_pdf(self) -> str:
        """
        Extrai texto de PDF usando pdfplumber, com fallback para PyPDF2.

        Returns:
            Texto bruto concatenado de todas as páginas.
        """
        text_parts = []

        # Tentativa primária: pdfplumber
        try:
            with pdfplumber.open(self.file_path) as pdf:
                for page_num, page in enumerate(pdf.pages, start=1):
                    page_text = page.extract_text() or ""
                    if page_text.strip():
                        text_parts.append(page_text)
                    logger.debug(f"Página {page_num}: {len(page_text)} caracteres")

            full_text = "\n\n".join(text_parts)

            # Se pdfplumber retornou texto suficiente, usa esse resultado
            if len(full_text.strip()) > 50:
                logger.info(f"PDF extraído via pdfplumber ({len(pdf.pages)} páginas)")
                return full_text

            logger.warning("pdfplumber retornou pouco texto, tentando PyPDF2 como fallback")

        except Exception as e:
            logger.warning(f"pdfplumber falhou ({e}), usando PyPDF2 como fallback")

        # Fallback: PyPDF2
        try:
            reader = PdfReader(str(self.file_path))
            text_parts = []
            for page_num, page in enumerate(reader.pages, start=1):
                page_text = page.extract_text() or ""
                text_parts.append(page_text)
                logger.debug(f"PyPDF2 - Página {page_num}: {len(page_text)} caracteres")

            full_text = "\n\n".join(text_parts)
            logger.info(f"PDF extraído via PyPDF2 ({len(reader.pages)} páginas)")
            return full_text

        except Exception as e:
            logger.exception(f"Ambos pdfplumber e PyPDF2 falharam: {e}")
            raise DocumentParseError(f"Falha total ao extrair PDF: {e}") from e

    def _parse_docx(self) -> str:
        """
        Extrai texto de arquivo Word (.docx).

        Inclui o conteúdo de parágrafos e tabelas.

        Returns:
            Texto bruto extraído.
        """
        document = Document(str(self.file_path))
        text_parts = []

        # Parágrafos do corpo
        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text)

        # Conteúdo de tabelas (importante em editais com cronogramas/valores)
        for table in document.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells)
                if row_text.strip(" |"):
                    text_parts.append(row_text)

        logger.info(
            f"DOCX extraído: {len(document.paragraphs)} parágrafos, "
            f"{len(document.tables)} tabelas"
        )
        return "\n".join(text_parts)

    def _parse_txt(self) -> str:
        """
        Lê arquivo de texto plano, tentando múltiplos encodings.

        Returns:
            Conteúdo bruto do arquivo.
        """
        encodings = ["utf-8", "latin-1", "cp1252", "iso-8859-1"]

        for encoding in encodings:
            try:
                with open(self.file_path, "r", encoding=encoding) as f:
                    content = f.read()
                logger.info(f"TXT lido com encoding {encoding}")
                return content
            except UnicodeDecodeError:
                continue

        raise DocumentParseError(
            f"Não foi possível decodificar o arquivo TXT com nenhum encoding testado"
        )


def parse_document(file_path: Union[str, Path]) -> str:
    """
    Função de conveniência para parsear um documento em uma chamada.

    Args:
        file_path: Caminho para o arquivo.

    Returns:
        Texto extraído e limpo.
    """
    parser = DocumentParser(file_path)
    return parser.parse()
